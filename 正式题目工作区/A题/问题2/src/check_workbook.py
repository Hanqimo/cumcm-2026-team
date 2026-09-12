"""Check every output cell against the selected run, plus template structure."""
from pathlib import Path
import hashlib,json,zipfile
import numpy as np
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
p=ROOT/'results/result2.xlsx'
with zipfile.ZipFile(p) as z:assert z.testzip() is None
w=load_workbook(p,read_only=True,data_only=False)
assert w.sheetnames==['温度','水分浓度']
x=np.loadtxt(ROOT/'runs/R006_N12800_s1024/solution_samples.csv',delimiter=',',skiprows=1)
report={'package_valid':True,'sheets':{},'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
for name,cols in [('温度',slice(1,22)),('水分浓度',slice(22,43))]:
    s=w[name]
    rows=s.iter_rows();header=next(rows)
    assert len(header)==22 and header[0].value=='时间\\到药材中心的距离'
    assert np.allclose([c.value for c in header[1:]],np.arange(21)*.1,rtol=0,atol=1e-14)
    maxdiff=0
    for i,row in enumerate(rows,1):
        assert len(row)==22
        assert row[0].value==i
        for j,c in enumerate(row[1:]):
            assert c.data_type=='n' and isinstance(c.value,(int,float)),(name,i,j,c.value)
            assert c.number_format=='0.0000',(name,i,j,c.number_format)
            expected=x[i,cols][j];diff=abs(c.value-expected);maxdiff=max(maxdiff,diff)
            assert diff<2e-12,(name,i,j,diff)
            assert f'{c.value:.4f}'==f'{expected:.4f}',(name,i,j,c.value,expected)
    assert i==10800
    report['sheets'][name]={'time_rows':10800,'radii':21,'values_checked':10800*21,'maximum_cell_difference':maxdiff,'four_decimal_render_values_agree':True}
report['all_cells_pass']=True
(ROOT/'verification/workbook_check.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False))
